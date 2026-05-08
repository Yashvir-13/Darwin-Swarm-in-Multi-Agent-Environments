from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
LOG_DIR = ROOT / "logs"
OUTPUT_PATH = ROOT / "Darwin_Swarm_Showcase.pptx"

SUMMARY_PATH = LOG_DIR / "summary.json"
FITNESS_PLOT = LOG_DIR / "fitness_curve.png"
BEHAVIOR_PLOT = LOG_DIR / "behavior_metrics.png"
BASELINE_PLOT = LOG_DIR / "baseline_comparison.png"
HEATMAP_PLOT = LOG_DIR / "best_team_heatmap.png"


BG = RGBColor(248, 249, 251)
NAVY = RGBColor(18, 40, 76)
TEAL = RGBColor(23, 128, 120)
GREEN = RGBColor(43, 122, 76)
ORANGE = RGBColor(214, 117, 61)
RED = RGBColor(170, 53, 53)
GRAY = RGBColor(90, 98, 112)
LIGHT = RGBColor(230, 235, 242)
WHITE = RGBColor(255, 255, 255)


def load_summary() -> dict:
    with open(SUMMARY_PATH, "r", encoding="utf-8") as handle:
        return json.load(handle)


def style_slide(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.7))
    text_frame = title_box.text_frame
    paragraph = text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = title
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = NAVY

    if subtitle:
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.65), Inches(0.98), Inches(11.5), Inches(0.4)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_paragraph = subtitle_frame.paragraphs[0]
        subtitle_run = subtitle_paragraph.add_run()
        subtitle_run.text = subtitle
        subtitle_run.font.size = Pt(12)
        subtitle_run.font.color.rgb = GRAY


def add_bullets(slide, left: float, top: float, width: float, height: float, bullets: list[str]) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(19)
        paragraph.font.color.rgb = NAVY
        paragraph.space_after = Pt(12)


def add_metric_card(slide, left: float, top: float, width: float, height: float, title: str, value: str, accent: RGBColor) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LIGHT

    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(0.16),
        Inches(height),
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()

    title_box = slide.shapes.add_textbox(
        Inches(left + 0.28), Inches(top + 0.18), Inches(width - 0.35), Inches(0.3)
    )
    title_frame = title_box.text_frame
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.text = title
    title_paragraph.font.size = Pt(12)
    title_paragraph.font.color.rgb = GRAY

    value_box = slide.shapes.add_textbox(
        Inches(left + 0.28), Inches(top + 0.48), Inches(width - 0.35), Inches(0.5)
    )
    value_frame = value_box.text_frame
    value_paragraph = value_frame.paragraphs[0]
    value_paragraph.text = value
    value_paragraph.font.size = Pt(24)
    value_paragraph.font.bold = True
    value_paragraph.font.color.rgb = NAVY


def add_image(slide, path: Path, left: float, top: float, width: float, height: float | None = None) -> None:
    if height is None:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    else:
        slide.shapes.add_picture(
            str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height)
        )


def add_caption(slide, text: str, left: float, top: float, width: float) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.4))
    frame = box.text_frame
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(11)
    run.font.color.rgb = GRAY


def add_architecture_boxes(slide) -> None:
    boxes = [
        (0.8, 2.15, 2.0, 1.1, "Input\n18 features", TEAL),
        (3.25, 2.15, 2.0, 1.1, "Hidden 1\n64 ReLU", GREEN),
        (5.7, 2.15, 2.0, 1.1, "Hidden 2\n64 ReLU", ORANGE),
        (8.15, 2.15, 2.0, 1.1, "Output\n5 actions", RED),
    ]

    for left, top, width, height, label, color in boxes:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = color
        shape.line.width = Pt(2)

        box = slide.shapes.add_textbox(
            Inches(left + 0.1), Inches(top + 0.22), Inches(width - 0.2), Inches(height - 0.2)
        )
        frame = box.text_frame
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = label
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = NAVY

    connectors = [(2.8, 2.7, 3.25, 2.7), (5.25, 2.7, 5.7, 2.7), (7.7, 2.7, 8.15, 2.7)]
    for x1, y1, x2, y2 in connectors:
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        line.line.color.rgb = GRAY
        line.line.width = Pt(2)


def add_pipeline(slide) -> None:
    steps = [
        ("Initialize", 0.7, TEAL),
        ("Evaluate", 2.55, GREEN),
        ("Select", 4.4, ORANGE),
        ("Crossover", 6.25, RED),
        ("Mutate", 8.1, TEAL),
        ("Repeat", 9.95, GREEN),
    ]
    y = 2.35
    for label, left, color in steps:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(y),
            Inches(1.45),
            Inches(0.9),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = color
        shape.line.width = Pt(2)

        box = slide.shapes.add_textbox(Inches(left), Inches(y + 0.22), Inches(1.45), Inches(0.3))
        para = box.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = label
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = NAVY

    for left in [2.15, 4.0, 5.85, 7.7, 9.55]:
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(left), Inches(2.8), Inches(left + 0.35), Inches(2.8)
        )
        line.line.color.rgb = GRAY
        line.line.width = Pt(2)


def build_presentation() -> Path:
    summary = load_summary()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    evolved = summary["baseline_results"]["Evolved"]
    random_team = summary["baseline_results"]["Random"]
    still_team = summary["baseline_results"]["Still"]
    config = summary["config"]
    final_gen = summary["final_generation"]

    def new_slide():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        style_slide(slide)
        return slide

    slide = new_slide()
    add_title(
        slide,
        "Darwin-Swarm",
        "Evolutionary Multi-Agent Coordination Sandbox",
    )
    add_bullets(
        slide,
        0.8,
        1.55,
        5.4,
        3.2,
        [
            "Neuroevolution instead of backpropagation or RL",
            "Agents do not learn during lifetime",
            "Population improves through selection, crossover, and mutation",
            "Current result: clear improvement from random behavior with an honest local-optimum limitation",
        ],
    )
    add_metric_card(slide, 0.9, 5.15, 2.0, 1.1, "Environment", "simple_spread", TEAL)
    add_metric_card(slide, 3.1, 5.15, 2.0, 1.1, "Agents", "3", GREEN)
    add_metric_card(slide, 5.3, 5.15, 2.0, 1.1, "Generations", str(config["num_generations"]), ORANGE)
    add_metric_card(slide, 7.5, 5.15, 2.0, 1.1, "Population", str(config["population_size"]), RED)
    add_image(slide, FITNESS_PLOT, 8.8, 1.45, 4.0, 3.0)
    add_caption(slide, "Fitness improves sharply from random initialization.", 8.75, 4.55, 4.1)

    slide = new_slide()
    add_title(slide, "Problem And Motivation", "Why evolve behavior instead of training it directly?")
    add_bullets(
        slide,
        0.8,
        1.45,
        5.8,
        4.6,
        [
            "Most AI systems rely on supervised learning or reinforcement learning.",
            "Those methods inject strong objectives and human priors.",
            "Darwin-Swarm asks whether structured multi-agent behavior can emerge from selection pressure alone.",
            "The project focuses on coordination, collision avoidance, and landmark coverage in a shared environment.",
        ],
    )
    quote = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(7.0),
        Inches(1.7),
        Inches(5.4),
        Inches(2.4),
    )
    quote.fill.solid()
    quote.fill.fore_color.rgb = WHITE
    quote.line.color.rgb = LIGHT
    qbox = slide.shapes.add_textbox(Inches(7.25), Inches(2.0), Inches(4.9), Inches(1.8))
    para = qbox.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = (
        "Can coordinated multi-agent behavior emerge without explicit lifetime learning?"
    )
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = NAVY
    add_bullets(
        slide,
        7.1,
        4.45,
        5.0,
        1.6,
        [
            "No external dataset is used.",
            "Experience is generated online through simulation rollouts.",
        ],
    )

    slide = new_slide()
    add_title(slide, "Technical Setup", "Exact configuration used in the current project")
    add_architecture_boxes(slide)
    add_bullets(
        slide,
        0.8,
        3.8,
        5.9,
        2.4,
        [
            "Input: 18 features per agent observation",
            "Outputs: 5 continuous action values",
            "Genome size: 5701 parameters per agent",
            "Two hidden layers, each with 64 ReLU units",
        ],
    )
    add_bullets(
        slide,
        6.8,
        3.8,
        5.3,
        2.4,
        [
            "Observation = self velocity, self position, landmark-relative positions, other-agent-relative positions, communication slots",
            "Communication slots exist but are zero in the current silent-agent setup",
            "Dataset: none; all experience comes from environment interaction",
        ],
    )

    slide = new_slide()
    add_title(slide, "Evolutionary Workflow", "How a generation is created and evaluated")
    add_pipeline(slide)
    add_bullets(
        slide,
        0.9,
        4.05,
        11.4,
        1.9,
        [
            "Each genome is evaluated in seeded team episodes.",
            "Fitness comes from environment reward, while behavior metrics track collisions, coverage, spread, and landmark distance.",
            "Top-performing genomes survive; new children are produced by crossover plus Gaussian mutation.",
        ],
    )

    slide = new_slide()
    add_title(slide, "Training Progress", "What changed over 160 generations")
    add_image(slide, FITNESS_PLOT, 0.7, 1.3, 6.1, 4.7)
    add_metric_card(slide, 7.2, 1.5, 2.2, 1.0, "Best Fitness", f"{summary['best_overall_fitness']:.2f}", GREEN)
    add_metric_card(slide, 9.6, 1.5, 2.2, 1.0, "Final Avg", f"{final_gen['avg_fitness']:.2f}", TEAL)
    add_metric_card(slide, 7.2, 2.75, 2.2, 1.0, "Final Coverage", f"{final_gen['coverage_rate']:.3f}", ORANGE)
    add_metric_card(slide, 9.6, 2.75, 2.2, 1.0, "Final Collisions", f"{final_gen['collision_rate']:.4f}", RED)
    add_bullets(
        slide,
        7.15,
        4.15,
        5.2,
        1.8,
        [
            "Average fitness improved dramatically from the random starting population.",
            "Improvement was rapid early, then plateaued into a stable local optimum.",
        ],
    )

    slide = new_slide()
    add_title(slide, "Behavior Metrics", "Evolution changed behavior, not only reward")
    add_image(slide, BEHAVIOR_PLOT, 0.7, 1.35, 7.2, 5.4)
    add_bullets(
        slide,
        8.2,
        1.8,
        4.5,
        3.9,
        [
            "Coverage increased substantially from the first generations.",
            "Collision rate fell by nearly an order of magnitude.",
            "Agent spread and landmark distance improved sharply early in training.",
            "Later generations stabilized rather than continuing to discover richer motion patterns.",
        ],
    )

    slide = new_slide()
    add_title(slide, "Baseline Comparison", "Evolved team vs simple baselines")
    add_image(slide, BASELINE_PLOT, 0.8, 1.45, 6.0, 4.8)
    add_bullets(
        slide,
        7.2,
        1.6,
        5.0,
        4.4,
        [
            f"Evolved reward: {evolved['team_reward_mean']:.1f}",
            f"Random reward: {random_team['team_reward_mean']:.1f}",
            f"Still reward: {still_team['team_reward_mean']:.1f}",
            "The evolved team clearly beats random behavior.",
            "However, it does not yet beat the still baseline on every metric, revealing a low-motion local optimum.",
        ],
    )

    slide = new_slide()
    add_title(slide, "Spatial Footprint", "Where the elite team spends time")
    add_image(slide, HEATMAP_PLOT, 0.7, 1.35, 7.1, 4.9)
    add_bullets(
        slide,
        8.1,
        1.7,
        4.5,
        4.1,
        [
            "Heatmaps show the learned team settling into safe regions with low collision behavior.",
            "This is visually useful for the showcase because it makes the learned strategy legible.",
            "The current weakness is also visible: agents may become too static rather than dynamically coordinating.",
        ],
    )

    slide = new_slide()
    add_title(slide, "Interpretation", "What we can honestly claim right now")
    add_bullets(
        slide,
        0.9,
        1.55,
        5.7,
        4.8,
        [
            "The project successfully demonstrates neuroevolution in a multi-agent environment.",
            "The population improves clearly from random initialization.",
            "Behavior becomes less chaotic and collisions drop significantly.",
            "The system produces reproducible artifacts for analysis and presentation.",
        ],
    )
    callout = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(7.0),
        Inches(1.8),
        Inches(5.3),
        Inches(2.1),
    )
    callout.fill.solid()
    callout.fill.fore_color.rgb = WHITE
    callout.line.color.rgb = ORANGE
    callout.line.width = Pt(2)
    tbox = slide.shapes.add_textbox(Inches(7.25), Inches(2.15), Inches(4.8), Inches(1.3))
    p = tbox.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Current limitation:\nlow-motion local optimum"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = NAVY
    add_bullets(
        slide,
        7.0,
        4.45,
        5.1,
        1.7,
        [
            "This is a meaningful result, not a failure.",
            "It gives the project a concrete next research direction.",
        ],
    )

    slide = new_slide()
    add_title(slide, "Next Steps", "How we would improve the next version")
    add_bullets(
        slide,
        0.9,
        1.6,
        5.5,
        4.5,
        [
            "Add a simple landmark-seeking scripted baseline",
            "Penalize inactivity or reward sustained landmark switching",
            "Track displacement and action magnitude to quantify collapse",
            "Run multi-seed experiments and compare robustness",
        ],
    )
    add_bullets(
        slide,
        6.8,
        1.6,
        5.0,
        4.5,
        [
            "Migrate fully to mpe2",
            "Sweep environment settings such as local_ratio and agent count",
            "Move to a richer environment like Waterworld after the analysis pipeline is stable",
            "Export GIFs or videos for side-by-side visual storytelling",
        ],
    )

    prs.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_presentation()
    print(path)
