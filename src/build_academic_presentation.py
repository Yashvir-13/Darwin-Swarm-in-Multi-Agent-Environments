from __future__ import annotations

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
LOG_DIR = ROOT / "logs"
OUTPUT_PATH = ROOT / "Darwin_Swarm_Academic_Format.pptx"

SUMMARY_PATH = LOG_DIR / "summary.json"
METRICS_CSV = LOG_DIR / "training_metrics.csv"
FITNESS_PLOT = LOG_DIR / "fitness_curve.png"
BEHAVIOR_PLOT = LOG_DIR / "behavior_metrics.png"
BASELINE_PLOT = LOG_DIR / "baseline_comparison.png"


BG = RGBColor(247, 249, 252)
NAVY = RGBColor(24, 41, 74)
TEAL = RGBColor(18, 133, 121)
GREEN = RGBColor(45, 122, 71)
ORANGE = RGBColor(214, 117, 61)
RED = RGBColor(176, 63, 63)
GRAY = RGBColor(92, 100, 112)
LIGHT = RGBColor(225, 231, 239)
WHITE = RGBColor(255, 255, 255)


TEAM_MEMBERS = [
    "Yashvir Singh - E23CSEU0293",
    "Shivesh Sharma - E23CSEU0271",
    "Rudraksh Srivastava - [Add Enrollment No.]",
    "Ayaan Khan - [Add Enrollment No.]",
]
SUPERVISOR = "[Add Supervisor / Guide Name]"
INSTITUTION = "[Add Institution / Department]"
PRESENTATION_DATE = "15 April 2026"


LITERATURE_REFERENCES = [
    "Stanley & Miikkulainen (2002) - NEAT",
    "Lehman & Stanley (2011) - Novelty Search",
    "Such et al. (2017) - Deep Neuroevolution",
    "Lowe et al. (2017) - MADDPG / MPE benchmark",
]


def load_summary() -> dict:
    with open(SUMMARY_PATH, "r", encoding="utf-8") as handle:
        return json.load(handle)


def load_metrics() -> tuple[dict[str, str], dict[str, str]]:
    with open(METRICS_CSV, "r", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    return rows[0], rows[-1]


def set_background(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.7))
    frame = box.text_frame
    para = frame.paragraphs[0]
    run = para.add_run()
    run.text = title
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = NAVY

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.95), Inches(11.8), Inches(0.35))
        sub_para = sub_box.text_frame.paragraphs[0]
        sub_run = sub_para.add_run()
        sub_run.text = subtitle
        sub_run.font.size = Pt(11)
        sub_run.font.color.rgb = GRAY


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    size: int = 18,
    bold: bool = False,
    color: RGBColor = NAVY,
    align=PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bullets(slide, left: float, top: float, width: float, height: float, bullets: list[str], size: int = 18) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, item in enumerate(bullets):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = item
        para.font.size = Pt(size)
        para.font.color.rgb = NAVY
        para.space_after = Pt(10)


def add_card(slide, left: float, top: float, width: float, height: float, title: str, body: str, accent: RGBColor) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LIGHT

    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(0.12),
        Inches(height),
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()

    add_textbox(slide, left + 0.2, top + 0.12, width - 0.28, 0.25, title, size=12, color=GRAY)
    add_textbox(slide, left + 0.2, top + 0.4, width - 0.28, height - 0.45, body, size=16, bold=True)


def add_image(slide, path: Path, left: float, top: float, width: float, height: float | None = None) -> None:
    if height is None:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    else:
        slide.shapes.add_picture(
            str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height)
        )


def add_reference_footer(slide, refs: list[str]) -> None:
    text = "References: " + " | ".join(refs)
    add_textbox(slide, 0.55, 7.02, 12.0, 0.28, text, size=9, color=GRAY)


def add_workflow(slide) -> None:
    labels = [
        ("Initialize\nPopulation", 0.65, TEAL),
        ("Run\nEpisodes", 2.65, GREEN),
        ("Compute\nFitness", 4.65, ORANGE),
        ("Select\nParents", 6.65, RED),
        ("Crossover +\nMutation", 8.65, TEAL),
        ("Next\nGeneration", 10.65, GREEN),
    ]
    y = 2.35
    for label, left, color in labels:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(y),
            Inches(1.55),
            Inches(1.0),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = color
        shape.line.width = Pt(2)
        add_textbox(slide, left + 0.08, y + 0.18, 1.39, 0.56, label, size=16, bold=True, align=PP_ALIGN.CENTER)

    for start in [2.2, 4.2, 6.2, 8.2, 10.2]:
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(start), Inches(2.85), Inches(start + 0.45), Inches(2.85)
        )
        line.line.color.rgb = GRAY
        line.line.width = Pt(2)


def build_presentation() -> Path:
    summary = load_summary()
    first_row, last_row = load_metrics()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    evolved = summary["baseline_results"]["Evolved"]
    random_team = summary["baseline_results"]["Random"]
    still_team = summary["baseline_results"]["Still"]
    config = summary["config"]

    def new_slide():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_background(slide)
        return slide

    slide = new_slide()
    add_title(slide, "Darwin-Swarm", "Evolving Emergent Behavior in Multi-Agent Sandboxes")
    add_textbox(slide, 0.8, 1.45, 11.5, 0.55, "Student Name(s) & Enrollment No.", size=16, bold=True, color=TEAL)
    add_bullets(slide, 1.0, 1.9, 5.4, 1.8, TEAM_MEMBERS, size=18)
    add_card(slide, 7.0, 1.7, 4.9, 1.0, "Supervisor / Guide", SUPERVISOR, ORANGE)
    add_card(slide, 7.0, 2.95, 4.9, 1.0, "Institution / Department", INSTITUTION, GREEN)
    add_card(slide, 7.0, 4.2, 4.9, 1.0, "Date", PRESENTATION_DATE, RED)
    add_textbox(
        slide,
        0.85,
        5.4,
        11.2,
        0.8,
        "A neuroevolution system that studies how multi-agent coordination can emerge without backpropagation or lifetime learning.",
        size=22,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    slide = new_slide()
    add_title(slide, "Abstract")
    add_bullets(
        slide,
        0.8,
        1.45,
        6.0,
        3.8,
        [
            "Research problem: Can coordination emerge in a multi-agent system without reinforcement learning or gradient-based updates during lifetime?",
            "Objective: Build a reproducible neuroevolution framework where neural-network-controlled agents evolve inside a shared environment.",
            "Methodology: Evaluate genomes in the simple_spread environment, rank by fitness, and generate new populations through selection, crossover, and Gaussian mutation.",
            f"Key result 1: Average fitness improved from {float(first_row['avg_fitness']):.1f} to {float(last_row['avg_fitness']):.1f}.",
            f"Key result 2: Collision rate dropped from {float(first_row['collision_rate']):.4f} to {float(last_row['collision_rate']):.4f}, though a low-motion local optimum remains.",
        ],
        size=17,
    )
    add_image(slide, FITNESS_PLOT, 7.15, 1.55, 5.2, 3.9)
    add_textbox(slide, 7.3, 5.55, 4.9, 0.45, "Current finding: evolution improves behavior clearly, but not yet beyond every baseline.", size=11, color=GRAY, align=PP_ALIGN.CENTER)

    slide = new_slide()
    add_title(slide, "Introduction To The Project", "Background, problem statement, and objectives")
    add_bullets(
        slide,
        0.8,
        1.45,
        5.8,
        4.8,
        [
            "Background: Most modern AI systems rely on supervised learning or reinforcement learning, which impose explicit objectives and strong human priors.",
            "Problem statement: Truly emergent collective behavior is difficult to study when reward shaping directly prescribes what agents should do.",
            "Project objective: Explore whether coordination can arise through evolutionary pressure alone in a shared multi-agent environment.",
            "Secondary objective: Measure how fitness, collisions, coverage, and agent spread change across generations.",
        ],
        size=18,
    )
    add_card(slide, 7.05, 1.7, 5.0, 1.0, "Research Question", "Can a population of fixed neural controllers evolve coordinated behavior without lifetime learning?", TEAL)
    add_card(slide, 7.05, 3.0, 5.0, 1.0, "System Type", "Simulation-driven neuroevolution, not supervised learning and not RL training", ORANGE)
    add_card(slide, 7.05, 4.3, 5.0, 1.0, "Data Source", "No external dataset; all experience is generated online through environment rollouts", GREEN)

    slide = new_slide()
    add_title(slide, "Introduction To The Project", "Methodology overview, tools, and data source")
    add_workflow(slide)
    add_bullets(
        slide,
        0.9,
        4.05,
        5.8,
        2.0,
        [
            "Environment: simple_spread_v3 from MPE / PettingZoo, with code preferring mpe2.",
            "Controller: 18-input, 2 hidden-layer feedforward neural network with 5 outputs.",
            "Data source: generated through seeded simulation episodes; no external dataset is used.",
        ],
        size=17,
    )
    add_bullets(
        slide,
        6.95,
        4.05,
        5.1,
        2.0,
        [
            "Tools: Python, NumPy, PyTorch, PettingZoo / MPE2, Matplotlib.",
            "Outputs: checkpoints, fitness curves, behavior plots, baseline comparisons, and heatmaps.",
        ],
        size=17,
    )

    slide = new_slide()
    add_title(slide, "Literature Review", "Foundational work in neuroevolution and open-ended search")
    add_card(
        slide,
        0.75,
        1.55,
        3.85,
        2.0,
        "Stanley & Miikkulainen (2002) - NEAT",
        "Established a foundational neuroevolution method that evolves neural network structure and weights together.",
        TEAL,
    )
    add_card(
        slide,
        4.75,
        1.55,
        3.85,
        2.0,
        "Lehman & Stanley (2011) - Novelty Search",
        "Argued that objective-driven search can become deceptive and that exploratory evolutionary search can reveal richer behaviors.",
        ORANGE,
    )
    add_card(
        slide,
        8.75,
        1.55,
        3.85,
        2.0,
        "Such et al. (2017) - Deep Neuroevolution",
        "Showed that genetic algorithms can remain competitive even when evolving large neural networks for control tasks.",
        GREEN,
    )
    add_bullets(
        slide,
        0.95,
        4.05,
        11.2,
        1.7,
        [
            "Relevance to this project: Darwin-Swarm inherits the idea that useful control policies can be discovered through evolution instead of backpropagation.",
            "Difference: our project focuses on multi-agent coordination rather than single-agent benchmark control.",
        ],
        size=17,
    )
    add_reference_footer(
        slide,
        [
            "Stanley & Miikkulainen (2002)",
            "Lehman & Stanley (2011)",
            "Such et al. (2017)",
        ],
    )

    slide = new_slide()
    add_title(slide, "Literature Review", "Multi-agent benchmarks and relevance to Darwin-Swarm")
    add_card(
        slide,
        0.8,
        1.55,
        5.2,
        1.8,
        "Lowe et al. (2017) - MADDPG / Multi-Agent Actor-Critic",
        "Highlighted the non-stationarity and coordination challenges in multi-agent learning and popularized the MPE benchmark family used widely for cooperative and competitive tasks.",
        RED,
    )
    add_card(
        slide,
        6.4,
        1.55,
        5.8,
        1.8,
        "MPE2 simple_spread environment",
        "Provides a compact benchmark where agents must cover landmarks while avoiding collisions, making it a good testbed for measuring emergent coordination.",
        TEAL,
    )
    add_bullets(
        slide,
        0.95,
        3.9,
        11.2,
        2.0,
        [
            "Literature gap addressed here: Many works optimize multi-agent policies through RL, whereas Darwin-Swarm studies the same kind of interaction space through evolutionary search.",
            "Project relevance: The benchmark is simple enough to visualize clearly yet rich enough to reveal coordination, collapse, or local-optimum behavior.",
        ],
        size=17,
    )
    add_reference_footer(slide, ["Lowe et al. (2017)", "MPE2 simple_spread official docs"])

    slide = new_slide()
    add_title(slide, "Methodologies", "Data collection method and experiment design")
    add_bullets(
        slide,
        0.8,
        1.45,
        6.0,
        4.8,
        [
            "Data collection method: simulation-based experiments, not surveys and not external tabular/image datasets.",
            "Each generation evaluates genomes through multiple seeded episodes in a shared environment.",
            "Fitness is computed from the agent's environment reward.",
            "Behavior metrics recorded: collision rate, landmark coverage rate, mean pairwise distance, and mean landmark distance.",
            "Baseline teams used: random-action team and still / zero-action team.",
        ],
        size=18,
    )
    add_image(slide, BASELINE_PLOT, 7.1, 1.55, 5.0, 3.8)
    add_textbox(slide, 7.25, 5.5, 4.8, 0.5, "Baselines help judge whether evolution is learning useful behavior beyond random exploration.", size=11, color=GRAY, align=PP_ALIGN.CENTER)

    slide = new_slide()
    add_title(slide, "Methodologies", "Tools, technologies, and model configuration")
    add_card(slide, 0.8, 1.55, 3.0, 1.1, "Programming Stack", "Python, NumPy, PyTorch", TEAL)
    add_card(slide, 4.05, 1.55, 3.0, 1.1, "Environment", "MPE2 / PettingZoo simple_spread_v3", GREEN)
    add_card(slide, 7.3, 1.55, 3.0, 1.1, "Visualization", "Matplotlib + heatmaps", ORANGE)
    add_card(slide, 10.55, 1.55, 2.0, 1.1, "Checkpoints", "PyTorch .pt files", RED)
    add_bullets(
        slide,
        0.9,
        3.15,
        5.8,
        2.8,
        [
            "Neural network architecture: input 18 -> hidden 64 -> hidden 64 -> output 5",
            "Activation function: ReLU on both hidden layers",
            "Genome representation: flattened vector of 5701 trainable parameters",
            "Learning during episode: none; weights are fixed during rollout",
        ],
        size=17,
    )
    add_bullets(
        slide,
        6.95,
        3.15,
        5.2,
        2.8,
        [
            "Evolution operators: rank-based selection, uniform crossover, Gaussian mutation",
            f"Population size in latest run: {config['population_size']}",
            f"Generations in latest run: {config['num_generations']}",
            f"Episodes per genome: {config['episodes_per_genome']}",
        ],
        size=17,
    )

    slide = new_slide()
    add_title(slide, "Methodologies", "Dataset description, features, and preprocessing")
    add_bullets(
        slide,
        0.85,
        1.45,
        5.8,
        4.9,
        [
            "External dataset: none",
            "Effective data source: environment observations generated online during episodes",
            "Observation size per agent: 18",
            "Action size per agent: 5",
            "Global state size: 54",
            "Agents: 3, Landmarks: 3",
        ],
        size=18,
    )
    add_bullets(
        slide,
        6.95,
        1.45,
        5.2,
        4.9,
        [
            "Observation features:",
            "1. Self velocity (2 values)",
            "2. Self position (2 values)",
            "3. Relative positions of 3 landmarks (6 values)",
            "4. Relative positions of 2 other agents (4 values)",
            "5. Communication slots from other agents (4 values, effectively zero in the current silent-agent setup)",
            "Preprocessing: observations converted to float tensors; output actions clipped to [0, 1]",
        ],
        size=16,
    )

    slide = new_slide()
    add_title(slide, "Methodologies", "Workflow diagram and implementation steps")
    add_workflow(slide)
    add_image(slide, BEHAVIOR_PLOT, 0.75, 4.0, 5.8, 2.45)
    add_bullets(
        slide,
        6.95,
        4.0,
        5.0,
        2.4,
        [
            "Step 1: initialize a random population",
            "Step 2: evaluate genomes in team episodes",
            "Step 3: compute fitness and behavior metrics",
            "Step 4: retain elites and selected parents",
            "Step 5: create offspring via crossover + mutation",
            "Step 6: repeat and save artifacts each run",
        ],
        size=16,
    )

    prs.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_presentation()
    print(path)
